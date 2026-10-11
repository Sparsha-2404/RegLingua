"""
Extract M1's labelled examples (SEBI/RBI/DPDP sheets) from
RegLingua_M1_FINAL_v2.xlsx into data/raw/labeled_examples_m1.csv.

This regenerates a file that docs/INTEGRATION_NOTES.md already documents as
existing (516 rows), but which was not actually included in this zip.
Re-run this any time the xlsx is updated -- it's a pure function of the xlsx,
nothing here should ever be hand-edited.

Usage:
    python3 scripts/extract_m1_labeled_csv.py
"""
from __future__ import annotations

import csv
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parents[1]
XLSX_PATH = REPO_ROOT / "RegLingua_M1_FINAL_v2.xlsx"
OUT_PATH = REPO_ROOT / "data" / "raw" / "labeled_examples_m1.csv"

SHEETS = ["SEBI", "RBI", "DPDP"]

# Column order in the CSV. `text` and `risk_level` are the two columns
# services/classifier/dataset.py actually requires; the rest are kept for
# traceability (source/citation) and for anyone doing error analysis later.
FIELDNAMES = [
    "id",
    "text",
    "language",
    "regulator",
    "topic",
    "risk_level",
    "source",
    "source_reference",
    "source_id",
    "document_title",
    "section",
    "effective_date",
    "source_url",
]


def extract() -> int:
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    rows_out = []

    for sheet_name in SHEETS:
        ws = wb[sheet_name]
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        col = {name: idx for idx, name in enumerate(header) if name is not None}

        for row in ws.iter_rows(min_row=2, values_only=True):
            row_id = row[col["ID"]]
            if row_id is None:
                continue  # trailing blank rows
            rows_out.append(
                {
                    "id": row_id,
                    "text": (row[col["TEXT"]] or "").strip(),
                    "language": row[col["LANGUAGE"]],
                    "regulator": row[col["REGULATOR"]],
                    "topic": row[col["TOPIC"]],
                    # Uppercase here so the CSV already matches the
                    # RiskLevel enum (LOW/MEDIUM/HIGH) without relying on
                    # dataset.py's runtime .upper() call.
                    "risk_level": (row[col["RISK LEVEL"]] or "").strip().upper(),
                    "source": row[col["SOURCE"]],
                    "source_reference": row[col["SOURCE REFERENCE"]],
                    "source_id": row[col["SOURCE_ID"]],
                    "document_title": row[col["DOCUMENT_TITLE"]],
                    "section": row[col["SECTION"]],
                    "effective_date": row[col["EFFECTIVE_DATE"]],
                    "source_url": row[col["SOURCE_URL"]],
                }
            )

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(OUT_PATH, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDNAMES)
        writer.writeheader()
        writer.writerows(rows_out)

    return len(rows_out)


if __name__ == "__main__":
    n = extract()
    print(f"Wrote {n} rows to {OUT_PATH}")
