"""
DAY 4 DELIVERABLE (M2): Create regulator/topic evaluation labels.

"M2: Create regulator/topic evaluation labels." (Day 4 plan)
Gate: "Query retrieves source-linked regulatory chunks."

M3 builds the embedding index + metadata filtering today; M2's job is to
produce the ground-truth labels that let retrieval actually be *measured*
(Day 5 gate: "EvidenceAgent returns selected chunks with source IDs";
Day 5's M2 task: "Evaluate retrieval queries against known relevant
sources").

Source: RegLingua_M1_FINAL.xlsx (SEBI/RBI/DPDP sheets, 516 real rows),
the same source scripts/extract_m1_labeled_csv.py already reads for Day 3.
This script is independent of that one and does NOT touch or overwrite
data/raw/labeled_examples_m1.csv -- it writes a new, separate evaluation
file used only for retrieval evaluation, not classifier training.

Usage:
    python3 scripts/build_regulator_topic_eval_labels.py
"""
from __future__ import annotations

import csv
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parents[1]
XLSX_PATH = REPO_ROOT / "data" / "evaluation" / "RegLingua_M1_FINAL.xlsx"
OUT_PATH = REPO_ROOT / "data" / "evaluation" / "regulator_topic_eval_m2.csv"

SHEETS = ["SEBI", "RBI", "DPDP"]

FIELDNAMES = [
    "query_id",
    "original_id",
    "text",
    "language",
    "true_regulator",
    "true_topic",
]


def build() -> int:
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    rows_out = []
    query_counter = 1

    for sheet_name in SHEETS:
        ws = wb[sheet_name]
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        col = {name: idx for idx, name in enumerate(header) if name is not None}

        for row in ws.iter_rows(min_row=2, values_only=True):
            original_id = row[col["ID"]]
            if original_id is None:
                continue  # trailing blank rows

            text = (row[col["TEXT"]] or "").strip()
            regulator = (row[col["REGULATOR"]] or "").strip().upper()
            topic = (row[col["TOPIC"]] or "").strip()
            language = row[col["LANGUAGE"]]

            if not text or not regulator:
                continue  # skip incomplete rows rather than crash the whole build

            rows_out.append(
                {
                    "query_id": f"EVALQ{query_counter:04d}",
                    "original_id": original_id,
                    "text": text,
                    "language": language,
                    "true_regulator": regulator,
                    "true_topic": topic,
                }
            )
            query_counter += 1

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(OUT_PATH, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDNAMES)
        writer.writeheader()
        writer.writerows(rows_out)

    return len(rows_out)


if __name__ == "__main__":
    n = build()
    print(f"Wrote {n} rows to {OUT_PATH}")
