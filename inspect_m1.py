import openpyxl
wb = openpyxl.load_workbook('RegLingua_M1_FINAL_v2.xlsx', read_only=True)
for s in wb.sheetnames:
    ws = wb[s]
    print('---', s, '---', ws.dimensions)
    row1 = next(ws.iter_rows(max_row=1))
    print([c.value for c in row1])
